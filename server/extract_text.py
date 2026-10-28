#!/usr/bin/env python3
"""Extract plain text from uploaded documents. Prints text to stdout."""
import os
import sys


def pdf_text(path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts = []
    for page in reader.pages[:400]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def docx_text(path):
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def xlsx_text(path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets:
        parts.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                parts.append("\t".join(values))
    return "\n".join(parts)


def pptx_text(path):
    from pptx import Presentation

    parts = []
    presentation = Presentation(path)
    for index, slide in enumerate(presentation.slides, 1):
        parts.append(f"# 第 {index} 页")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def main():
    if len(sys.argv) < 2:
        print("")
        return 0
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            text = pdf_text(path)
        elif ext == ".docx":
            text = docx_text(path)
        elif ext == ".xlsx":
            text = xlsx_text(path)
        elif ext == ".pptx":
            text = pptx_text(path)
        else:
            text = ""
    except Exception as exc:  # Missing optional library or damaged file.
        sys.stderr.write(str(exc))
        print("")
        return 0
    sys.stdout.write(text[:400000])
    return 0


if __name__ == "__main__":
    sys.exit(main())
